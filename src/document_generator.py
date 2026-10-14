import os
import zipfile
import tempfile
import shutil
from docx import Document

class CanaryDocumentGenerator:
    def __init__(self, listener_url: str):
        self.listener_url = listener_url

    def create_word_canary(self, output_path: str, title: str = "Confidential Executive Payroll & Credentials"):
        """Generate an MS Word document embedded with an HTTP tracking bug."""
        doc = Document()
        doc.add_heading(title, level=1)
        doc.add_paragraph("RESTRICTED INTERNAL DOCUMENT — STRICTLY CONFIDENTIAL.")
        doc.add_paragraph("Contains cloud infrastructure credentials, internal hostnames, and database endpoints.")
        
        temp_dir = tempfile.mkdtemp()
        temp_docx = os.path.join(temp_dir, "base.docx")
        doc.save(temp_docx)

        unzip_dir = os.path.join(temp_dir, "unpacked")
        with zipfile.ZipFile(temp_docx, 'r') as zip_ref:
            zip_ref.extractall(unzip_dir)

        doc_xml_path = os.path.join(unzip_dir, "word", "document.xml")
        if os.path.exists(doc_xml_path):
            with open(doc_xml_path, 'r', encoding='utf-8') as f:
                content = f.read()

            tracking_pixel = f'<w:p><w:r><w:fldChar w:fldCharType="begin"/><w:instrText>INCLUDEPICTURE "{self.listener_url}/canary/word_doc_triggered" \\d</w:instrText><w:fldChar w:fldCharType="separate"/><w:fldChar w:fldCharType="end"/></w:r></w:p>'
            
            if "</w:body>" in content:
                content = content.replace("</w:body>", f"{tracking_pixel}</w:body>")

            with open(doc_xml_path, 'w', encoding='utf-8') as f:
                f.write(content)

        shutil.make_archive(output_path.replace('.docx', ''), 'zip', unzip_dir)
        final_zip = output_path.replace('.docx', '') + '.zip'
        if os.path.exists(output_path):
            os.remove(output_path)
        os.rename(final_zip, output_path)

        shutil.rmtree(temp_dir)
        return output_path
